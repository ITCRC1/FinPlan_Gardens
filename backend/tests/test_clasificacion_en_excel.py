# -*- coding: utf-8 -*-
"""La clasificación fijo/variable se baja a Excel, se llena y se vuelve a subir.

Owner, 2026-09-09: *«necesito que revises la configuración, la forma de asignar
el % de fijo o variable. Veo esa asignación muy complicada, debe ser muy fácil.
Inclusive que se baje a Excel y ahí se haga la asignación y se vuelva a subir;
veo que en la pantalla es muy difícil»*.

Y enseguida, sobre el alcance: *«el criterio no debe ser por mes»* · *«debe ser
completo, uno solo sin diferencial mes»*.

## Qué se vigila acá

Lo que hace peligroso un viaje redondo por archivo es que **el archivo vuelve
sin que nadie mire lo que trae**. Las tres formas de perder plata en silencio:

1. Una celda **vacía** tratada como cero — un descuido al filtrar volvería 100%
   fija una cuenta variable, y el punto de equilibrio se movería sin que nada lo
   diga.
2. El apareo **por posición** en vez de por `id` — una columna corrida y cada
   porcentaje aterriza en la cuenta de al lado.
3. Un valor imposible **aceptado a la fuerza** (recortado a 0 o a 100) en vez de
   rechazado con su motivo.

Las tres se prueban acá. La cuarta —que el archivo se aplique sin que nadie lo
haya visto— la cubre `aplicar=False` por omisión en el endpoint.
"""
import io
import json
import pathlib

import pytest
from openpyxl import load_workbook

from app.export.break_even_xlsx import (COL_PCT, FILA_ENCABEZADO,
                                        construir_plantilla, leer_plantilla)

RAIZ = pathlib.Path(__file__).resolve().parents[1]


def _filas():
    return [
        {"id": "r1", "departamento": "Rooms", "be_section": "OPERATING EXPENSES",
         "dept_code": "0110", "account": "7310", "account_name": "Laundry",
         "pl_line": "OPEX_ROOMS", "pct_variable": 1.0, "amount": 12345.67,
         "original_class": "Variable", "excluded_from_be": False},
        {"id": "r2", "departamento": "A&G", "be_section": "OVERHEAD",
         "dept_code": "0180", "account": "7160", "account_name": "Electricity",
         "pl_line": "OH_ADMIN", "pct_variable": 0.0, "amount": 8000.0,
         "original_class": "Fixed Cost", "excluded_from_be": False},
        {"id": "r3", "departamento": "Property", "be_section": "BELOW GOP",
         "dept_code": "0250", "account": "8060", "account_name": "Income tax",
         "pl_line": "INCOME_TAXES", "pct_variable": 0.0, "amount": 900.0,
         "original_class": "Fixed Cost", "excluded_from_be": True},
    ]


# ── El archivo que baja ──────────────────────────────────────────────────────

def test_baja_una_fila_por_regla_y_el_id_viaja():
    """El `id` es lo que hace posible el apareo al volver. Sin él habría que
    aparear por (departamento, cuenta), que es justo lo que el usuario puede
    reordenar en Excel."""
    wb = load_workbook(io.BytesIO(construir_plantilla(
        _filas(), "Amarena", "BUDGET Final 2026")))
    ws = wb.active
    ids = [ws.cell(FILA_ENCABEZADO + 1 + i, 1).value for i in range(3)]
    assert ids == ["r1", "r2", "r3"]


def test_el_porcentaje_baja_en_ESCALA_de_0_a_100():
    """En la base es una fracción (0,4). En el archivo va 40.

    Quien llena el archivo escribe «40», no «0.4»: pedirle la fracción es pedirle
    que traduzca, y una traducción a mano en 798 filas es un error esperando.
    """
    wb = load_workbook(io.BytesIO(construir_plantilla(
        _filas(), "Amarena", "BUDGET Final 2026")))
    ws = wb.active
    assert ws.cell(FILA_ENCABEZADO + 1, COL_PCT).value == 100.0   # 1.0
    assert ws.cell(FILA_ENCABEZADO + 2, COL_PCT).value == 0.0     # 0.0


def test_el_monto_viaja_para_no_clasificar_a_ciegas():
    """Da lo mismo equivocarse en una cuenta de $12.345 que en una de $12 —y no
    da lo mismo. Es la misma razón por la que la pantalla lo trae."""
    wb = load_workbook(io.BytesIO(construir_plantilla(
        _filas(), "Amarena", "BUDGET Final 2026")))
    ws = wb.active
    montos = [ws.cell(FILA_ENCABEZADO + 1 + i, 10).value for i in range(3)]
    assert montos == [12345.67, 8000.0, 900.0]


def test_la_hoja_va_protegida_y_solo_el_porcentaje_se_edita():
    """⚠️ No es desconfianza: el archivo vuelve y se aparea por `id`. Una
    columna corrida a mano rompe el apareo, y el que sube no se entera."""
    wb = load_workbook(io.BytesIO(construir_plantilla(
        _filas(), "Amarena", "BUDGET Final 2026")))
    ws = wb.active
    assert ws.protection.sheet is True
    editable = ws.cell(FILA_ENCABEZADO + 1, COL_PCT)
    bloqueada = ws.cell(FILA_ENCABEZADO + 1, 5)          # la cuenta
    assert editable.protection.locked is False
    assert bloqueada.protection.locked is not False


def test_la_linea_excluida_no_se_ofrece_para_editar():
    """El impuesto de renta no tiene parte variable que asignar. Dejarlo
    editable sería ofrecer una decisión que el motor ignora."""
    wb = load_workbook(io.BytesIO(construir_plantilla(
        _filas(), "Amarena", "BUDGET Final 2026")))
    ws = wb.active
    celda = ws.cell(FILA_ENCABEZADO + 3, COL_PCT)        # r3, excluida
    assert celda.value is None
    assert celda.protection.locked is True


def test_el_fijo_es_una_FORMULA_y_no_un_numero():
    """Dos columnas editables que tienen que sumar 100 son dos formas de decir
    lo mismo y una oportunidad de que discrepen. Si el archivo volviera diciendo
    30 y 30, habría que elegir a cuál creerle."""
    wb = load_workbook(io.BytesIO(construir_plantilla(
        _filas(), "Amarena", "BUDGET Final 2026")))
    ws = wb.active
    fijo = ws.cell(FILA_ENCABEZADO + 1, COL_PCT + 1).value
    assert isinstance(fijo, str) and fijo.startswith("=100-")


# ── El archivo que vuelve ────────────────────────────────────────────────────

def _con_valores(valores: dict):
    """Baja la plantilla, escribe los valores dados y devuelve los bytes."""
    crudo = construir_plantilla(_filas(), "Amarena", "BUDGET Final 2026")
    wb = load_workbook(io.BytesIO(crudo))
    ws = wb.active
    for i, rid in enumerate(["r1", "r2", "r3"]):
        if rid in valores:
            ws.cell(FILA_ENCABEZADO + 1 + i, COL_PCT).value = valores[rid]
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_el_viaje_redondo_conserva_el_id_y_el_valor():
    leido = dict(leer_plantilla(_con_valores({"r1": 40, "r2": 25})))
    assert leido["r1"] == 40
    assert leido["r2"] == 25


def test_una_celda_VACIA_vuelve_como_vacia_y_no_como_cero():
    """⚠️ La trampa que este viaje redondo hace posible.

    Vacío significa «no toqué esta fila». Convertirlo en cero volvería 100% fija
    una cuenta variable por un descuido —borrar una celda, filtrar y no darse
    cuenta— y el punto de equilibrio se movería sin que nada avise. Para poner
    cero hay que escribir `0`, y eso también se comprueba.
    """
    leido = dict(leer_plantilla(_con_valores({"r1": None, "r2": 0})))
    assert leido["r1"] is None, "una celda vacía llegó como un valor"
    assert leido["r2"] == 0, "un cero escrito a mano tiene que llegar como cero"


def test_el_lector_NO_decide_nada_de_negocio():
    """Devuelve el valor CRUDO. No convierte, no valida, no recorta.

    El criterio —qué es un cero, qué está fuera de rango, qué fila está
    excluida— lo aplica el endpoint, que es el único que ve la base. Un lector
    que además decide es un lugar más donde el criterio puede diferir.
    """
    leido = dict(leer_plantilla(_con_valores({"r1": 150, "r2": "abc"})))
    assert leido["r1"] == 150          # fuera de rango, pero lo devuelve igual
    assert leido["r2"] == "abc"        # ni siquiera es un número


def test_el_encabezado_se_encuentra_aunque_la_leyenda_crezca():
    """Se busca por CONTENIDO, no por número de fila: si no, cada cambio de
    formato invalidaría en silencio el archivo que alguien tenía a medio
    llenar."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "una leyenda"
    ws["A2"] = "mucho mas larga"
    ws["A3"] = "que la original"
    ws.cell(4, 1, "id")
    ws.cell(4, 2, "% Variable")
    ws.cell(5, 1, "r9")
    ws.cell(5, 2, 33)
    buf = io.BytesIO()
    wb.save(buf)
    assert leer_plantilla(buf.getvalue()) == [("r9", 33)]


# ── El endpoint ──────────────────────────────────────────────────────────────

def _fuente() -> str:
    import inspect

    from app.api import break_even_api
    return inspect.getsource(break_even_api.subir_plantilla)


def test_sube_EN_SECO_por_omision():
    """⚠️ Mover el % de una cuenta cambia el punto de equilibrio. El que sube
    tiene derecho a ver la lista antes de que se escriba."""
    from app.api import break_even_api

    firma = __import__("inspect").signature(break_even_api.subir_plantilla)
    assert firma.parameters["aplicar"].default.default is False


def test_lo_que_no_se_puede_aplicar_se_RECHAZA_y_se_dice_cual():
    """No se aplica a medias ni se recorta al rango. Cada rechazo lleva su
    motivo, porque un rechazo sin motivo se lee como un error del sistema."""
    f = _fuente()
    for motivo in ("no es un número", "fuera de 0 a 100",
                   "es una línea excluida del cálculo",
                   "la fila aparece dos veces en el archivo"):
        assert motivo in f, motivo


def test_se_avisa_lo_que_el_archivo_NO_trajo():
    """Se puede subir un archivo filtrado, y está bien. Pero el que sube cree
    que subió todo, así que se le dice cuántas reglas no venían."""
    assert "no_venian_en_el_archivo" in _fuente()


def test_el_apareo_es_por_id_y_nunca_por_posicion():
    f = _fuente()
    assert "reglas.get(rid)" in f
    assert "no es una regla de esta propiedad" in f


def test_la_plantilla_es_de_la_propiedad_ENTERA_y_sin_mes():
    """Owner: *«el criterio no debe ser por mes; debe ser completo, uno solo sin
    diferencial mes»*.

    El modelo ya era así —`be_cost_classification` no tiene columna de mes— y el
    endpoint no acepta uno: el monto que muestra es el del año completo
    (`montos_del_escenario(..., 0)`), que es la escala a la que se decide si un
    gasto sigue a la venta.
    """
    import inspect

    from app.api import break_even_api

    firma = inspect.signature(break_even_api.bajar_plantilla)
    assert "month" not in firma.parameters, (
        "la plantilla volvió a aceptar un mes: la clasificación no tiene mes")
    f = inspect.getsource(break_even_api.bajar_plantilla)
    assert "montos_del_escenario(db, s, 0)" in f
    # Y no filtra por departamento: baja la propiedad entera.
    assert "dept_slug" not in firma.parameters


def test_el_modelo_no_tiene_columna_de_mes():
    """La garantía de fondo: aunque alguien agregue un selector de mes a la
    pantalla, no hay dónde guardar una clasificación distinta por mes."""
    from app.models.break_even import BeCostClassification

    columnas = {c.name for c in BeCostClassification.__table__.columns}
    assert not (columnas & {"month", "mes", "period", "periodo"}), columnas
