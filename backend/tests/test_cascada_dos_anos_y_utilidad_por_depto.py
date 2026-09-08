# -*- coding: utf-8 -*-
"""
EL EXCEL DEL P&L DETAIL: LOS DOS AÑOS, Y LA UTILIDAD POR DEPARTAMENTO.

Owner, 2026-09-07: *«que salgan los 2 años comparativos 12 meses en el excel»* ·
*«tab de cierre pon la vista de profit … por departamento»*.

Dos defectos que arreglan estas pruebas:

**1 · La hoja Cascada bajaba una sola versión.** Usaba `series[0]` y nada más,
así que las versiones que el owner elegía para comparar en pantalla no llegaban
al archivo: el Excel mostraba un año donde la pantalla mostraba dos.

**2 · El cuadro de Cierre no abría la utilidad por departamento.** Traía sólo los
diez totales.

⚠️ **Y la trampa que hay que vigilar:** los nombres de departamento se repiten en
las TRES secciones de la cascada — «Rooms» está en REVENUES, en Operating
Expenses y en Operating Profit. Seleccionar por rótulo devuelve cualquiera de las
tres, y devolvía **distinto en cada lado**: el exportador indexaba con un dict
(se queda con el ÚLTIMO) y la pantalla con `find` (toma el PRIMERO). Por eso el
corte es por SECCIÓN y por posición.
"""
import io

import openpyxl
import pytest

from app.export.pl_detail_excel import (
    SECCION_UTILIDAD, _filas_de_seccion, export_pl_detail,
)

#: Valores distintos por sección para que una confusión se note en la prueba.
ING_ROOMS, GASTO_ROOMS, UTIL_ROOMS = 1000.0, 400.0, 600.0
ING_FB, GASTO_FB, UTIL_FB = 500.0, 300.0, 200.0


def _doce(v: float) -> list[float]:
    return [v] * 12


def _filas() -> list[dict]:
    """Una cascada mínima con las tres secciones y el rótulo «Rooms» repetido."""
    def det(rot, v1, v2):
        return {"tipo": "det", "rotulo": rot,
                "series": [_doce(v1), _doce(v2)]}
    return [
        {"tipo": "sec", "rotulo": "REVENUES", "series": [None, None]},
        det("Rooms", ING_ROOMS, ING_ROOMS * 2),
        det("F&B", ING_FB, ING_FB * 2),
        {"tipo": "tot", "rotulo": "TOTAL REVENUES",
         "series": [_doce(ING_ROOMS + ING_FB), _doce((ING_ROOMS + ING_FB) * 2)]},
        {"tipo": "esp", "rotulo": "", "series": [None, None]},
        {"tipo": "sec", "rotulo": "Operating Expenses", "series": [None, None]},
        det("Rooms", GASTO_ROOMS, GASTO_ROOMS * 2),
        det("F&B", GASTO_FB, GASTO_FB * 2),
        {"tipo": "tot", "rotulo": "Total Operationg expenses",
         "series": [_doce(GASTO_ROOMS + GASTO_FB), _doce((GASTO_ROOMS + GASTO_FB) * 2)]},
        {"tipo": "esp", "rotulo": "", "series": [None, None]},
        {"tipo": "sec", "rotulo": SECCION_UTILIDAD, "series": [None, None]},
        det("Rooms", UTIL_ROOMS, UTIL_ROOMS * 2),
        det("F&B", UTIL_FB, UTIL_FB * 2),
        {"tipo": "tot", "rotulo": "OPERATING PROFIT",
         "series": [_doce(UTIL_ROOMS + UTIL_FB), _doce((UTIL_ROOMS + UTIL_FB) * 2)]},
    ]


def _kpis() -> dict:
    return {"rooms_available": _doce(100), "rooms_occupied": _doce(50),
            "guests": _doce(90), "rooms_revenue": _doce(ING_ROOMS)}


def _datos() -> dict:
    return {
        "titulo_ambito": "Consolidado", "nota_ambito": "Hotel + Club",
        "escenario": "ACTUAL Final 2026", "year": 2026,
        "versiones": [
            {"scenario_id": "a", "escenario": "ACTUAL Final 2026", "kpis": _kpis()},
            {"scenario_id": "b", "escenario": "ACTUAL Final 2025", "kpis": _kpis()},
        ],
        "filas": _filas(),
        "clave": ["TOTAL REVENUES", "OPERATING PROFIT"],
        "clases_rotulos": [("payroll", "Total Payroll and Benefits")],
        "clases": [{"payroll": _doce(10)}, {"payroll": _doce(20)}],
        "control": {"ingresos": 0.0, "gastos": 0.0, "utilidad": 0.0,
                    "diferencia": 0.0},
    }


def _libro():
    return openpyxl.load_workbook(io.BytesIO(export_pl_detail(_datos(), 7)))


# ── El corte por seccion ─────────────────────────────────────────────────────

def test_la_seccion_devuelve_la_utilidad_y_no_el_ingreso():
    """La prueba de la trampa: «Rooms» esta tres veces y hay que traer la de
    utilidad, no la de ingreso ni la de gasto."""
    filas = _filas_de_seccion(_filas(), SECCION_UTILIDAD)
    assert [f["rotulo"] for f in filas] == ["Rooms", "F&B"]
    assert filas[0]["series"][0][0] == UTIL_ROOMS, "trajo otra seccion"
    assert filas[1]["series"][0][0] == UTIL_FB


def test_la_seccion_corta_en_el_total_y_no_sigue():
    """No se lleva las filas de la seccion siguiente."""
    filas = _filas_de_seccion(_filas(), "REVENUES")
    assert [f["series"][0][0] for f in filas] == [ING_ROOMS, ING_FB]


def test_una_seccion_que_no_existe_no_revienta():
    assert _filas_de_seccion(_filas(), "NO EXISTE") == []


# ── La hoja Cascada: los dos anos ────────────────────────────────────────────

def test_la_cascada_trae_las_dos_versiones():
    ws = _libro()["Cascada"]
    textos = [c.value for fila in ws.iter_rows(max_row=6) for c in fila]
    assert "ACTUAL Final 2026" in textos
    assert "ACTUAL Final 2025" in textos, "la segunda version no bajo"


def test_la_cascada_da_trece_columnas_por_version():
    """Doce meses mas el Full Year, un bloque por version.

    La fila de la cabecera se BUSCA, no se fija: `_titulo` puede crecer y una
    prueba con el numero escrito se rompe sin que nada este mal.
    """
    ws = _libro()["Cascada"]
    cab = next(([c.value for c in r] for r in ws.iter_rows(max_row=10)
                if any(c.value == "Ene" for c in r)), None)
    assert cab is not None, "no se encontro el piso de los meses"
    assert cab.count("Ene") == 2, f"esperaba dos bloques, cabecera={cab[:30]}"
    assert cab.count("Full Year") == 2


def test_el_full_year_de_cada_version_es_la_suma_de_sus_doce_meses():
    ws = _libro()["Cascada"]
    fila = next(r for r in ws.iter_rows(min_row=7)
                if (r[0].value or "").strip() == "TOTAL REVENUES")
    v1 = [c.value for c in fila[1:14]]
    v2 = [c.value for c in fila[14:27]]
    assert v1[12] == pytest.approx(sum(v1[:12]))
    assert v2[12] == pytest.approx(sum(v2[:12]))
    # Y la segunda version es el doble, como la armo el fixture: si los bloques
    # se pisaran, los dos darian lo mismo.
    assert v2[12] == pytest.approx(v1[12] * 2)


# ── La hoja Cierre: la utilidad por departamento ─────────────────────────────

def test_el_cierre_abre_la_utilidad_por_departamento():
    ws = _libro()["Cierre"]
    rotulos = [(c.value or "").strip() for c in ws["A"] if isinstance(c.value, str)]
    assert f"{SECCION_UTILIDAD} — por departamento" in rotulos
    assert "Rooms" in rotulos and "F&B" in rotulos


def test_el_cierre_muestra_la_utilidad_y_no_el_ingreso_de_ese_departamento():
    """Si alguien vuelve a seleccionar por rotulo, esta prueba lo agarra."""
    ws = _libro()["Cierre"]
    fila = None
    visto_encabezado = False
    for r in ws.iter_rows():
        rot = (r[0].value or "").strip() if isinstance(r[0].value, str) else ""
        if rot == f"{SECCION_UTILIDAD} — por departamento":
            visto_encabezado = True
            continue
        if visto_encabezado and rot == "Rooms":
            fila = r
            break
    assert fila is not None, "no se encontro el renglon de Rooms bajo la utilidad"
    # Primer bloque = el mes 7, una sola celda: la utilidad, no el ingreso.
    assert fila[1].value == pytest.approx(UTIL_ROOMS)
